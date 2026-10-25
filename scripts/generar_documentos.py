from pathlib import Path
from datetime import datetime, timezone
import json, re, sys
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "docs"
OUT.mkdir(exist_ok=True)
DATE = "23 de octubre de 2026"
raw = json.loads((ROOT / "reports/vitest-results.json").read_text(encoding="utf-8-sig"))
cases = []
for suite in raw["testResults"]:
    filename = Path(suite["name"]).relative_to(ROOT).as_posix()
    for case in suite["assertionResults"]:
        match = re.search(r"(?:UT|UI|CK|CP)\d+", case["fullName"])
        cases.append({"id": match.group() if match else case["title"], "nombre": case["fullName"], "archivo": filename, "estado": case["status"], "errores": case["failureMessages"]})
cases.sort(key=lambda c: ({"UT": 0, "UI": 1, "CK": 2, "CP": 3}.get(c["id"][:2], 4), int(re.search(r"\d+", c["id"]).group()) if re.search(r"\d+", c["id"]) else 0))
results = {"fecha": datetime.fromtimestamp(raw["startTime"] / 1000, timezone.utc).isoformat(), "resultado": "passed" if raw["success"] else "failed", "total": raw["numTotalTests"], "aprobadas": raw["numPassedTests"], "fallidas": raw["numFailedTests"], "archivos": len(raw["testResults"]), "pruebas": cases}
coverage = json.loads((ROOT / "coverage/coverage-summary.json").read_text(encoding="utf-8-sig"))
browser_path = ROOT / "reports/browser-results.json"
if not browser_path.exists():
    browser_path = OUT / "evidencia/navegador.json"
browser = json.loads(browser_path.read_text(encoding="utf-8-sig"))
total = coverage["total"]
evidence = OUT / "evidencia"
evidence.mkdir(exist_ok=True)
(evidence / "vitest.json").write_text(json.dumps(results, ensure_ascii=False, indent=2), encoding="utf-8")
(evidence / "cobertura.json").write_text(json.dumps(total, ensure_ascii=False, indent=2), encoding="utf-8")
P = lambda s: ("p", s)
H = lambda s: ("h", s)
T = lambda headers, rows, widths=None: ("t", (headers, rows, widths))
B = lambda s: ("b", s)

def table(doc, headers, rows, widths):
    t = doc.add_table(rows=1, cols=len(headers))
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.autofit = False
    if widths:
        for col, width in zip(t.columns, widths):
            col.width = Inches(width)
    for c, text in zip(t.rows[0].cells, headers):
        c.text = str(text)
    for row in rows:
        for c, text in zip(t.add_row().cells, row):
            c.text = str(text)
    for i, row in enumerate(t.rows):
        trpr = row._tr.get_or_add_trPr()
        trpr.append(OxmlElement("w:cantSplit"))
        if i == 0:
            trpr.append(OxmlElement("w:tblHeader"))
        for j, c in enumerate(row.cells):
            if widths: c.width = Inches(widths[j])
            c.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            tcpr = c._tc.get_or_add_tcPr()
            borders = OxmlElement("w:tcBorders")
            for edge in ("top","left","bottom","right"):
                el = OxmlElement("w:" + edge)
                for key,value in {"val":"single","sz":"5","color":"D9D9D9"}.items(): el.set(qn("w:"+key),value)
                borders.append(el)
            tcpr.append(borders)
            shd = OxmlElement("w:shd")
            shd.set(qn("w:fill"), "243348" if i == 0 else ("F2F5F8" if i % 2 == 0 else "FFFFFF"))
            tcpr.append(shd)
            margins = OxmlElement("w:tcMar")
            for edge in ("top","left","bottom","right"):
                el = OxmlElement("w:"+edge); el.set(qn("w:w"),"85"); el.set(qn("w:type"),"dxa"); margins.append(el)
            tcpr.append(margins)
            for p in c.paragraphs:
                p.paragraph_format.space_after = Pt(0)
                p.paragraph_format.line_spacing = 1.08
                for r in p.runs:
                    r.font.size = Pt(9.5)
                    r.font.color.rgb = RGBColor.from_string("FFFFFF" if i == 0 else "20252B")
                    r.bold = i == 0


def make(name, title, subtitle, pages):
    if sys.argv[1:] and name not in sys.argv[1:]: return
    d = Document()
    sec=d.sections[0]
    sec.page_width=Inches(8.5);sec.page_height=Inches(11)
    sec.top_margin=sec.bottom_margin=Inches(.8)
    sec.left_margin=sec.right_margin=Inches(1)
    normal=d.styles["Normal"]
    normal.font.name="Calibri";normal.font.size=Pt(11);normal.font.color.rgb=RGBColor(0,0,0)
    normal.paragraph_format.space_after=Pt(8)
    normal.paragraph_format.line_spacing=1.12
    for sn,size in [("Title",26),("Subtitle",12),("Heading 1",17),("Heading 2",12)]:
        st=d.styles[sn];st.font.name="Calibri";st.font.size=Pt(size);st.font.color.rgb=RGBColor(0,0,0)
        st.paragraph_format.space_before=Pt(12 if sn.startswith("Heading") else 0)
        st.paragraph_format.space_after=Pt(8)
    d.add_paragraph(title,"Title");d.add_paragraph(subtitle,"Subtitle")
    d.add_paragraph("MiTienda  |  Full Stack II  |  Versión 2  |  " + DATE)
    footer=sec.footer.paragraphs[0];footer.alignment=2
    r=footer.add_run("MiTienda   ");r.font.size=Pt(9)
    f=OxmlElement("w:fldSimple");f.set(qn("w:instr"),"PAGE");footer._p.append(f)
    md=["# "+title,"",subtitle,"",DATE,""]
    for pi, blocks in enumerate(pages):

        for kind, data in blocks:
            if kind=="h":
                heading=d.add_paragraph(data,"Heading 1")
                if pi and blocks.index((kind,data)) == 0: heading.paragraph_format.page_break_before=True
                md+=["## "+data,""]
            elif kind=="b": d.add_paragraph(data,"Heading 2");md+=["### "+data,""]
            elif kind=="p": d.add_paragraph(data);md += [data,""]
            elif kind=="t":
                headers,rows,widths=data
                table(d,headers,rows,widths)
                md += ["| "+" | ".join(map(str,headers))+" |","| "+" | ".join(["---"]*len(headers))+" |"]+["| "+" | ".join(str(x).replace("|","/") for x in r)+" |" for r in rows]+[""]
    d.core_properties.title=title;d.core_properties.subject=subtitle;d.core_properties.author="Equipo MiTienda";d.core_properties.last_modified_by="Equipo MiTienda";d.core_properties.comments=""
    for node in list(d.styles.element.xpath(".//w:pBdr")) + list(d.element.xpath(".//w:pBdr")):
        node.getparent().remove(node)
    d.save(OUT/(name+".docx"))
    (OUT/(name+".md")).write_text("\n".join(md),encoding="utf-8")
    print(name+".docx")

ers=[
[
H("1 Propósito y alcance"),
P("Esta especificación define el comportamiento de MiTienda para la evaluación de Full Stack II. La versión 2 consolida catálogo, usuarios, compra simulada y administración, e incorpora los cambios de experiencia, pruebas y reportes. Sirve al equipo y al docente para revisar qué se implementa y cómo se acepta."),
P("El producto es una aplicación React y Vite que se ejecuta en el navegador. Los datos se guardan en localStorage del mismo origen. La entrega cubre una demostración funcional local; no incluye servidor de autenticación, base de datos compartida, cobros bancarios ni entregas físicas."),
B("Actores"),
T(["Actor","Operaciones"],[["Visitante","Explorar, filtrar, leer el blog, registrar cuenta, usar el carrito y comprar como invitado."],["Cliente","Iniciar y cerrar sesión; usar el catálogo y la compra con datos de contacto precargados."],["Administrador","Gestionar catálogo, categorías y usuarios; revisar pedidos, estados, inventario y ventas."]],[1.25,5.25]),
B("Entorno y arquitectura"),
P("Node.js 24 o superior ejecuta las herramientas y Vitest. React presenta las vistas, React Router administra las rutas y los contextos comparten productos, categorías, carrito, pedidos y sesión. El CSS adapta la tienda a escritorio y móvil."),
P("La línea base del repositorio es 396c877. Los contratos de datos existentes se conservan y se añaden campos de entrega, costo de envío y pago simulado."),
B("Convenciones"),
P("RF identifica requisitos funcionales, RNF requisitos de calidad y RN reglas del negocio. Cada requisito incluye un resultado observable. La cobertura y los casos ejecutados se registran en el documento de testing.")
],
[
H("2 Catálogo y usuarios"),
B("RF01 Inicio organizado"),
P("La página inicial debe mostrar una cabecera con acceso a categorías, bloques de categorías y dos selecciones de hasta cuatro productos cada una. Un producto seleccionado no debe repetirse entre ambas selecciones. Aceptación: se muestran como máximo ocho tarjetas y enlaces al catálogo completo."),
B("RF02 Búsqueda y filtros"),
P("El usuario debe buscar por nombre, marca, descripción o código sin distinguir mayúsculas ni tildes. En catálogo y categorías podrá combinar marca, precio mínimo y máximo, disponibilidad y orden por precio o nombre. Aceptación: los resultados y el contador reflejan la combinación; una búsqueda sin resultados muestra una salida para limpiar filtros."),
B("RF03 Categorías y subcategorías"),
P("La navegación debe permitir explorar categorías principales y sus subcategorías. Una categoría principal incluye los productos de sus hijos directos. La categoría Ofertas conserva su selección de destacados y productos de las categorías promocionales."),
B("RF04 Ficha y carrito"),
P("Cada tarjeta debe mostrar nombre, imagen, marca, descripción, precio CLP y acción de agregar. La ficha presenta el detalle del producto. Un producto agotado no admite añadidos y las cantidades respetan el stock del catálogo."),
B("RF05 Registro de cliente"),
P("El formulario público debe registrar exclusivamente clientes con nombre, correo válido y contraseña de al menos cuatro caracteres. El correo se recorta y normaliza a minúsculas y no puede repetirse. El rol administrador solo se gestiona desde el panel."),
B("RF06 Sesión y acceso"),
P("El inicio de sesión debe validar las credenciales de la demo. El encabezado debe permitir abrirlo como diálogo con cierre por botón o Escape; /login permanece disponible. Un cliente o visitante que intente acceder a /admin será llevado al login."),
B("RF07 Información y contacto"),
P("Nosotros debe explicar el proyecto y su equipo; el blog debe ofrecer guías con detalle navegable. Contacto debe validar nombre, correo y mensaje, guardar una consulta local y confirmar expresamente que no se envió un correo.")
],
[
H("3 Compra y pago de demostración"),
B("RF08 Carrito persistente"),
P("El carrito debe permitir aumentar y disminuir cantidades, eliminar con confirmación y vaciar. Debe recalcular el total y conservar únicamente id y cantidad. Sin productos debe mostrar acceso al catálogo y ocultar la acción de pago."),
B("RF09 Entrega y contacto"),
P("El paso de entrega debe solicitar receptor, correo y celular chileno. El envío exige calle y número, comuna y región. El retiro no exige dirección. El usuario podrá volver entre pasos sin perder lo completado."),
B("RF10 Tarjeta de prueba"),
P("El paso de pago debe informar dos números de demostración: 4242 4242 4242 4242 para aprobación y 4000 0000 0000 0002 para rechazo. Exige titular ficticio, vencimiento futuro MM/AA y CVV de tres dígitos. Otros números no avanzan."),
B("RF11 Revisión y aprobación"),
P("Antes de confirmar se muestran receptor, entrega, últimos cuatro dígitos y total. Al aprobar, se crea un único pedido, se descuentan las unidades compradas y se vacía el carrito. El comprobante presenta identificador, artículos, total y entrega, y permite imprimir."),
B("RF12 Rechazo y reintento"),
P("La tarjeta de rechazo debe mostrar una vista de error y permitir corregir los datos. Un rechazo no crea pedido, no descuenta stock y no vacía el carrito. El reintento aprobado crea un solo pedido."),
B("RF13 Tratamiento de datos de pago"),
P("El número completo, titular de tarjeta, vencimiento y CVV permanecen solo en el estado del formulario. El pedido guardado no contiene esos campos. No se invoca una pasarela externa."),
T(["Regla","Condición"],[["RN01","El stock y las cantidades son enteros y no negativos; para comprar la cantidad es mayor a cero."],["RN02","Envío simulado de $4.990 si el subtotal es menor a $100.000; desde $100.000 es gratuito."],["RN03","Retiro sin costo; total de compra igual a subtotal más costo de envío."],["RN04","Un pedido cancelado repone el stock una vez y no admite nuevos cambios de estado."]],[.8,5.7])
],
[
H("4 Administración y reportes"),
B("RF14 Productos"),
P("El administrador debe listar, buscar por código o nombre, crear y editar productos. Los nuevos códigos deben ser únicos. El formulario conserva marca, descripción, categoría, precio, stock, umbral crítico, imagen y destacado; admite rutas locales o imágenes HTTPS."),
B("RF15 Reporte de inventario"),
P("La ruta /admin/products/reports debe mostrar cantidad de productos, unidades, productos críticos y valor del inventario a precio de venta. El filtro crítico incluye agotados y stock menor o igual al umbral. La descarga CSV contiene solo la selección visible y escapa separadores, comillas y prefijos de fórmula."),
B("RF16 Categorías"),
P("El panel debe crear, editar y eliminar categorías y subcategorías de dos niveles. La interfaz impide eliminar categorías con hijos o productos. Evita nombres repetidos en un mismo nivel y que una categoría con hijos pase a ser subcategoría."),
B("RF17 Usuarios"),
P("El panel debe listar, consultar, crear y editar usuarios y sus roles. El correo identifica al usuario y debe ser único. El detalle debe permitir consultar el historial de pedidos del usuario. El perfil muestra los datos del administrador y acceso a edición."),
B("RF18 Pedidos"),
P("El panel debe listar pedidos, filtrar por estado, consultar su detalle y crear pedidos manuales con productos disponibles. El detalle debe presentar datos del receptor y dirección cuando corresponda. Cambiar a cancelado devuelve el stock y bloquea cambios posteriores."),
B("RF19 Ventas y dashboard"),
P("El dashboard resume ingresos, productos, inventario, usuarios y pedidos recientes. El reporte de ventas permite agrupar por día, semana o mes, comparar períodos y desglosar por producto, marca y categoría. Los pedidos cancelados se excluyen de ventas."),
B("Rutas principales"),
T(["Área","Rutas"],[["Tienda","/ · /products · /products/:id · /categories · /categories/:id"],["Compra","/cart · /checkout · /payment como ruta compatible"],["Usuarios e información","/login · /signup · /about · /contact · /blog · /blog/:id"],["Panel","/admin y secciones products, categories, orders, users, reports y profile"]],[1.4,5.1])
],
[
H("5 Calidad y datos"),
T(["Requisito","Criterio verificable"],[["RNF01 Adaptación","La tienda debe funcionar a 390 px y en escritorio de 1440 px sin desbordamiento horizontal del documento."],["RNF02 Accesibilidad","Inputs etiquetados, foco visible, botones con nombre, enlace de salto, diálogo nativo y respeto a movimiento reducido."],["RNF03 Mantenibilidad","Componentes con una responsabilidad, contextos para estado compartido y reglas de compra como funciones puras."],["RNF04 Reproducibilidad","npm ci, npm test, npm run test:coverage y npm run build deben poder ejecutarse con Node 24."],["RNF05 Persistencia","Un recargo mantiene datos del mismo navegador; JSON no válido usa valores por defecto."],["RNF06 Alcance de seguridad","La sesión y el control de roles son demostrativos. No se deben usar credenciales o tarjetas reales."]],[1.45,5.05]),
B("Modelo de datos"),
T(["Entidad o clave","Datos relevantes"],[["productos","id, código, nombre, marca, precio, imagen, descripción, stock, stockCritico, categoriaId, destacado."],["categorias","id, nombre y padreId nulo para principales; identificador del padre para subcategorías."],["usuarios y usuarioLogueado","Nombre, correo, contraseña y rol de la demo; datos personales de ejemplo."],["carrito","Lista de id de producto y cantidad; los demás datos se leen del catálogo."],["pedidos","id, fecha, estado, items, subtotal, costoEnvio, total, nombre, usuario, teléfono, tipoEntrega y dirección."],["consultas","Identificador, fecha, nombre, correo, asunto y comentario guardados localmente."]],[1.6,4.9]),
P("Los pedidos conservan una copia del nombre y precio de sus artículos para preservar el detalle histórico. El catálogo y las categorías guardan una versión de sus datos iniciales; la migración de imágenes no reemplaza cambios de inventario realizados por el administrador.")
],
[
H("6 Aceptación y trazabilidad"),
P("El flujo principal se acepta al completar una compra con envío, observar su comprobante, encontrar el pedido en administración y verificar que el carrito quede vacío y el stock disminuya. El flujo alternativo se acepta al rechazar una tarjeta, conservar el estado y aprobar un reintento."),
T(["Requisitos","Casos Vitest","Verificación adicional"],[["RF01 a RF04","UT19 a UT23 y UI01 a UI07","Inicio, catálogo y categorías en navegador."],["RF05 y RF06","UI08 a UI11, UI25, UI33 y UI34","Login emergente y acceso de administrador."],["RF07","UI12 y UI18 a UI20","Rutas informativas y blog."],["RF08 a RF13","UT01 a UT18, UI15, UI16 y CK01 a CK08","Rechazo, reintento aprobado y revisión del pedido."],["RF14 a RF19","UT24, UT25, UT28 a UT31, UI13, UI14, UI17 y UI26 a UI37","Recorrido de panel y CSV de inventario."],["RNF01 a RNF05","UT26 y UT27; suite de rutas","Build y revisión a 390 y 1440 px."]],[1.1,2.6,2.8]),
B("Condiciones de entrega"),
P("La entrega debe incluir código en una rama de trabajo, lockfile, pruebas Vitest ejecutables, informe de cobertura con métricas y brechas, propuesta de pruebas y este ERS. El código se revisa localmente antes de publicar."),
B("Límites y evolución"),
P("La demo no ofrece sincronización entre dispositivos, autorización del lado del servidor, cifrado de contraseñas, correo real ni pagos o logística reales. Una versión productiva requiere backend, base de datos, autenticación y pasarela de pagos, además de pruebas de integración sobre esos servicios."),
P("No se establece un SLA, una certificación de accesibilidad ni un porcentaje obligatorio de cobertura porque no forman parte de la rúbrica disponible. Los resultados medidos y las pruebas pendientes se describen de manera separada en el informe de testing."),
B("Fuentes del proyecto"),
P("Repositorio base: https://github.com/BDecinti/Ecommerce-full-stack-2. Contratos y organización: README.md. Implementación y pruebas: src. Esta especificación documenta la versión local preparada para revisión el 23 de octubre de 2026.")
]
]
make("ERS_v2","Especificación de requisitos del software","Tienda de tecnología y gaming MiTienda",ers)

propuesta=[
[
H("1 Estrategia y alcance"),
P("La estrategia combina reglas unitarias, componentes React y recorridos integrados de la compra. Vitest ejecuta los casos y mide la cobertura con su proveedor V8; React Testing Library dispara eventos y consulta la interfaz por nombres accesibles. JSDOM provee el DOM simulado. La configuración se encuentra en vite.config.js y src/test/setup.js."),
P("Se implementaron "+str(results["total"])+" casos automatizados. Cada archivo .test.jsx o .test.js está junto al componente o módulo probado. Por ejemplo, ProductCard.jsx y ProductCard.test.jsx comparten src/components. Los 24 componentes reutilizables tienen un archivo de pruebas propio, incluidos los 9 del panel. src/App.test.jsx cubre rutas e integración; src/test solo contiene utilidades. Cada caso limpia el DOM y el almacenamiento."),
T(["Tipo exigido","Aplicación concreta","Casos"],[["Renderizado","Tarjeta, inicio, catálogo vacío, rutas y panel.","UI01, UI04, UI06, UI07, UI18 a UI32"],["Props","Producto, insignia y configuración inicial de filtros.","UI01 y UI05"],["Estado","Carrito, catálogo, formularios y pedido persistente.","UI02, UI03, UI11, UI14 a UI17"],["Eventos","Click, escritura, selección, envío y cambio de estado.","UI05, UI08, UI09, UI12, UI35 a UI37 y CK01 a CK08"]],[1.1,4.2,1.2]),
B("Pruebas de los componentes"),
P("CP01 a CP22 comprueban navegación, diálogos, cierre con Escape, foco, buscador, contador del carrito, avisos temporizados y carrusel. CP23 a CP34 cubren título, iconos, barra lateral, estados, datos del pedido, subtotales y cantidades limitadas al stock. CP35 a CP40 verifican filtros, límites de precio e imágenes con respaldo local ante errores."),
B("Organización del equipo"),
P("El README asigna catálogo a Alessandro, compra y usuarios a Nicolás y administración a Bruno. La revisión de estas pruebas puede repartirse por esas áreas."),
B("Ejecución reproducible"),
P("Con Node.js 24, ejecutar npm ci, npm test, npm run test:coverage y npm run build. Los resultados de casos se generan en reports/vitest-results.json y los porcentajes en coverage/coverage-summary.json. coverage/index.html permite inspeccionar líneas no ejercitadas.")
],
[
H("2 Casos prioritarios de compra"),
T(["Caso","Preparación y acción","Resultado esperado"],[["CK01","Entrar al pago con carrito vacío.","No aparece confirmación de pago; se ofrece catálogo."],["CK02 y CK03","Seleccionar envío y probar celular inválido.","Dirección requerida; costo visible; contacto inválido bloquea avance."],["CK04","Un producto de $10.000, envío y tarjeta 4242.","Pedido de $14.990, dirección registrada, stock -1 y carrito vacío."],["CK05","Mismo carrito con tarjeta 4000 terminada en 0002.","Rechazo; stock, carrito y pedidos sin cambios."],["CK06","Después del rechazo, cambiar a 4242 y confirmar.","Un pedido de retiro, sin cobro de envío."],["CK07","Ingresar otro número de tarjeta.","La revisión se bloquea con un error explicativo."],["CK08","Completar envío, avanzar y volver.","La dirección y los datos del receptor permanecen."]],[.9,3.1,2.5]),
B("Casos de límites y reglas"),
P("UT01 comprueba $99.999 y $100.000 para envío gratuito. UT03 a UT08 verifican los campos de entrega por separado. UT10 a UT15 cubren ambos números demo, formato y vencimiento. UT16 incluye cantidad cero, negativa, fraccionaria, superior al stock y producto inexistente."),
B("Integridad de administración"),
P("UI14 verifica crear, editar, descontar, reponer y eliminar en ProductContext. UI35 crea una categoría desde su formulario. UI36 rechaza un código duplicado. UI37 cancela un pedido, comprueba la reposición y verifica el bloqueo de cambios posteriores.")
],
[
H("3 Propuestas de ampliación"),
P("Los siguientes escenarios se proponen para una siguiente iteración. No se contabilizan como aprobados ni aumentan la cobertura informada hasta contar con su implementación y ejecución."),
T(["Prioridad","Prueba propuesta","Motivo"],[["Alta","Crear, editar y eliminar por interfaz todos los registros permitidos; verificar dependencias.","Ampliar la regresión del panel más allá de rutas, código duplicado y categoría nueva."],["Alta","Simular cuota de almacenamiento, datos con estructura inválida y dos pestañas editando.","La persistencia actual es local y las escrituras no constituyen una transacción de servidor."],["Alta","Cambiar stock o precio desde administración durante la revisión de compra.","Detectar cambios concurrentes y mantener el resumen coherente."],["Media","Navegación completa por teclado, lector de pantalla y contraste medido.","Complementar etiquetas, foco visible y diálogo nativo."],["Media","Ejecutar la misma compra en Firefox, Safari y Android e iOS reales.","La revisión de navegador ejecutada corresponde a Edge en Windows."],["Posterior","Integración con API, autenticación, correo y pasarela sandbox cuando exista backend.","Estos servicios no están implementados en esta entrega."]],[.8,3.4,2.3]),
B("Criterios de revisión"),
P("Toda corrección funcional debe acompañarse de un caso que reproduzca el fallo. La suite y el build deben pasar antes de integrar cambios. La cobertura se evalúa por riesgo y por líneas sin ejecutar; un porcentaje global no sustituye la revisión de los flujos."),
P("Se aplica la corrección del docente: Vitest reemplaza al ejecutor anterior y los casos se guardan junto a sus componentes. npm test ejecuta una vez; npm run test:watch repite al editar. Los componentes se verifican por su comportamiento visible, propiedades, estado y eventos.")
]
]
make("Propuesta_pruebas_Vitest","Propuesta de pruebas con Vitest","Estrategia y escenarios de validación de MiTienda",propuesta)

def pct(obj, metric): return str(obj[metric]["pct"]) + " %"
selected=["components/Header.jsx","components/ConfirmDialog.jsx","components/CartMessage.jsx","components/admin/OrderProductPicker.jsx","utils/checkout.js","pages/checkout/Payment.jsx","utils/filtrarProductos.js","components/ProductFilters.jsx","context/CartContext.jsx","context/ProductContext.jsx","context/AuthContext.jsx","utils/reporteVentas.js","pages/admin/ProductReports.jsx"]
files={}
for key,value in coverage.items():
    k=key.replace("\\","/")
    if "/src/" in k: files[k.split("/src/",1)[1]]=value
metrics=[[label,total[key]["covered"],total[key]["total"],pct(total,key)] for label,key in [("Líneas","lines"),("Sentencias","statements"),("Funciones","functions"),("Ramas","branches")]]
covpages=[
[
H("1 Resultado de la ejecución"),
P("La suite de "+str(results["archivos"])+" archivos terminó con "+str(results["aprobadas"])+" pruebas aprobadas de "+str(results["total"])+", sin fallos. Las pruebas cubren reglas del catálogo y compra, componentes React, estado compartido, navegación y regresión del panel. La revisión funcional anterior completó compra rechazada, reintento aprobado y descarga del inventario. La revisión del 23 de octubre verifica separación de acciones en escritorio y móvil y la vista HTML portátil."),
T(["Métrica global de src","Cubierto","Total","Porcentaje"],metrics,[2.7,1,1,1.8]),
P("La medición incluye todos los archivos JavaScript y JSX de src, aunque no se importen en un caso. Excluye CSS, fotografías, archivos .test y src/test. Se usa @vitest/coverage-v8, versión 5.0.5. Las métricas se calculan con el mapeo de código de Vitest; no son directamente comparables con las del ejecutor anterior."),
P("La cobertura describe la ejecución automatizada, no certifica que todo el sistema esté libre de errores. La revisión visual y los recorridos en Edge se informan por separado y no se suman a los porcentajes de Vitest."),
B("Entorno y reproducción"),
P("Windows, Node.js 24, Vitest 5.0.5, React Testing Library y JSDOM. Comandos: npm ci; npm test; npm run test:coverage; npm run build. Los scripts se encuentran en package.json, vite.config.js y src/test/setup.js."),
P("Los reportes de datos se generan en reports/vitest-results.json y coverage/coverage-summary.json. El reporte navegable se encuentra en coverage/index.html. La revisión corresponde a la versión local basada en 396c877.")
],
[
H("2 Cobertura de módulos relevantes"),
T(["Archivo de src","Líneas","Funciones","Ramas"],[[k,pct(files[k],"lines"),pct(files[k],"functions"),pct(files[k],"branches")] for k in selected],[3.6,.95,.95,1]),
B("Verificación en navegador"),
P("Edge en Windows recorrió "+str(len(browser["rutas"]))+" rutas de tienda y administración sin errores de JavaScript. Se comprobó búsqueda de G502, carrito, dirección, rechazo con conservación del carrito, reintento aprobado, pedido persistido y descarga CSV. Se revisó el inicio en 1440 por 1000 px y en 390 por 844 px; el documento móvil midió 390 px, sin desbordamiento."),
B("Brechas y alcance"),
P("Los porcentajes inferiores a 100 % señalan rutas o funciones que requieren más casos. La tabla de módulos no representa una meta aprobada por el docente. La propuesta complementaria detalla las ampliaciones sugeridas."),
P("La navegación del panel prueba que las vistas se presentan; no equivale a probar todas las combinaciones de sus formularios. No se ejecutaron pruebas sobre un backend, pagos bancarios, correo real, sincronización entre equipos, carga concurrente o certificación de accesibilidad."),
B("Registro de casos"),
P("Las páginas siguientes enumeran todos los casos ejecutados y sus estados. UT identifica reglas unitarias, UI componentes y navegación, CK compra integrada y CP los nuevos casos de componentes reutilizables. Los casos de compra ejercitan contextos reales y almacenamiento del navegador simulado.")
]
]
specs=results["pruebas"]
for offset in range(0,len(specs),20):
    rows=[]
    for case in specs[offset:offset+20]:
        name=case["nombre"].split(case["id"],1)[-1].strip()
        rows.append([case["id"],name,"Aprobada" if case["estado"]=="passed" else case["estado"]])
    covpages.append([H(str(3+offset//20)+" Casos ejecutados"),T(["ID","Verificación","Estado"],rows,[.65,5,.85])])
make("Cobertura_testing","Cobertura de testing","Resultados medidos y trazabilidad de las pruebas de MiTienda",covpages)
